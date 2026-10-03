#!/usr/bin/env python3
"""Loads returned data-request workbooks (docs/templates) into the platform. Dry run unless --apply is given.

    pip install openpyxl
    python scripts/import-workbook.py --token <ADMIN TOKEN> returned/*.xlsx            # check only
    python scripts/import-workbook.py --token <ADMIN TOKEN> --apply returned/*.xlsx    # load

Loads: outlet sheets (1A, 2A) -> customers (new ones added, existing ones updated),
3A Market size -> the outlet universe, 4A Products -> products. Other sheets are for analysis and are skipped.
"""
import argparse, csv, io, json, re, sys, urllib.request, urllib.error
from openpyxl import load_workbook

CUSTOMER_COLS = ["type", "name", "specialty", "segment", "territory", "parent", "phone", "email", "address", "city",
                 "latitude", "longitude", "target_visits_per_month", "channel", "outlet_class"]
PRODUCT_COLS = ["item_code", "name", "therapeutic_area", "standard_cost", "reorder_level"]
CLASS_ENUM = {"teachinghospital": "TeachingHospital", "regionalhospital": "RegionalHospital", "districthospital": "DistrictHospital",
              "retailpharmacychain": "PharmacyChain", "pharmacychain": "PharmacyChain", "independentpharmacy": "IndependentPharmacy",
              "otcshop": "OtcShop", "clinicother": "ClinicOrOther", "clinicorother": "ClinicOrOther", "clinic": "ClinicOrOther"}


def rows(ws, cols):
    """Rows from row 6 down; row 4 holds the headings (matched by name, so column order does not matter)."""
    heads = {str(c.value).strip(): i for i, c in enumerate(ws[4], start=0) if c.value}
    out = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        rec = {}
        for c in cols:
            v = r[heads[c]] if c in heads and heads[c] < len(r) else None
            if hasattr(v, "strftime"): v = v.strftime("%Y-%m-%d")
            rec[c] = "" if v is None else str(v).strip()
        if any(rec.values()): out.append(rec)
    return out


def to_csv(recs, cols):
    buf = io.StringIO(); w = csv.writer(buf); w.writerow(cols)
    for r in recs: w.writerow([r[c] for c in cols])
    return buf.getvalue()


def call(method, url, token, body, ctype):
    req = urllib.request.Request(url, data=body.encode("utf-8"), method=method,
                                 headers={"Authorization": f"Bearer {token}", "Content-Type": ctype})
    try:
        with urllib.request.urlopen(req, timeout=120) as r: return r.status, r.read().decode()
    except urllib.error.HTTPError as e: return e.code, e.read().decode()


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("files", nargs="+")
    p.add_argument("--token", required=True)
    p.add_argument("--api", default="http://localhost:5111")
    p.add_argument("--apply", action="store_true", help="really load the data (default is a dry run)")
    a = p.parse_args()
    api = a.api.rstrip("/") + "/api/v1"
    failed = False
    for f in a.files:
        wb = load_workbook(f, data_only=True)
        print(f"\n== {f}")
        for ws in wb.worksheets:
            name = ws.title
            if name.startswith("1A") or name.startswith("2A"):
                recs = rows(ws, CUSTOMER_COLS)
                if not recs: print(f"  {name}: empty, skipped"); continue
                s, t = call("POST", f"{api}/customers/import?dryRun={str(not a.apply).lower()}&onDuplicate=update", a.token, to_csv(recs, CUSTOMER_COLS), "text/csv")
                print(f"  {name}: {len(recs)} rows -> HTTP {s}")
                print("   ", t[:1500])
                failed |= s >= 300
            elif name.startswith("3A"):
                recs = rows(ws, ["region", "outlet_class", "outlets", "source"])
                body = []
                for r in recs:
                    cls = CLASS_ENUM.get(re.sub(r"[^a-z]", "", r["outlet_class"].lower()))
                    if not cls or not r["outlets"].replace(",", "").split(".")[0].isdigit():
                        print(f"  {name}: row skipped (check kind/number): {r}"); failed = True; continue
                    body.append({"region": r["region"] or None, "outletClass": cls, "outlets": int(r["outlets"].replace(",", "").split(".")[0]), "source": r["source"] or None})
                if not body: print(f"  {name}: empty, skipped"); continue
                if a.apply:
                    s, t = call("PUT", f"{api}/rtm/universe", a.token, json.dumps(body), "application/json")
                    print(f"  {name}: {len(body)} rows -> HTTP {s} {t[:300]}"); failed |= s >= 300
                else: print(f"  {name}: {len(body)} rows look valid (dry run)")
            elif name.startswith("4A"):
                recs = rows(ws, PRODUCT_COLS)
                if not recs: print(f"  {name}: empty, skipped"); continue
                if a.apply:
                    s, t = call("POST", f"{api}/erp/import/products", a.token, to_csv(recs, PRODUCT_COLS), "text/csv")
                    print(f"  {name}: {len(recs)} rows -> HTTP {s} {t[:300]}"); failed |= s >= 300
                else: print(f"  {name}: {len(recs)} rows look valid (dry run)")
            else:
                print(f"  {name}: analysis only, not imported")
    print("\nDry run only. Add --apply to load." if not a.apply else "\nDone.")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
